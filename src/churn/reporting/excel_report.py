"""Builds excel/telco_churn_report.xlsx.

Every number on the summary tabs is a live formula over the Data sheet, so the
workbook can be handed to someone who wants to tweak the campaign assumptions
without touching Python.
"""
from __future__ import annotations

from contextlib import closing
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from churn.config import ROOT, get_logger, load_config
from churn.pipeline.load import connect
from churn.pipeline.transform import TENURE_LABELS

log = get_logger(__name__)

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3B57")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
KPI_FILL = PatternFill("solid", fgColor="EEF4FC")
THIN = Side(style="thin", color="D9D8D3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PCT = "0.0%"
MONEY = "$#,##0;($#,##0);-"
MONEY2 = "$#,##0.00"
INT = "#,##0"

DATA_COLS = [
    "customer_id", "gender", "senior_citizen", "partner", "dependents",
    "contract_type", "payment_method", "internet_service", "tenure", "tenure_band",
    "n_services", "monthly_charges", "total_charges", "churn",
    "churn_probability", "risk_tier", "segment_name", "split",
]


def _font(bold=False, color="0B0B0B", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _header(ws, row: int, values: list[str], col: int = 1) -> None:
    for j, v in enumerate(values):
        c = ws.cell(row=row, column=col + j, value=v)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def _title(ws, text: str, sub: str | None = None) -> None:
    ws["A1"] = text
    ws["A1"].font = _font(bold=True, size=16)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = _font(color="52514E", italic=True)
    ws.sheet_view.showGridLines = False


def _load_frame(segments: pd.DataFrame | None) -> pd.DataFrame:
    with closing(connect()) as con:
        df = pd.read_sql_query("SELECT * FROM vw_customer_360", con)
    scores_path = ROOT / "data" / "processed" / "churn_scores.csv"
    if scores_path.exists():
        s = pd.read_csv(scores_path)[["customer_id", "churn_probability", "risk_tier", "split"]]
        df = df.merge(s, on="customer_id", how="left")
    else:
        df["churn_probability"], df["risk_tier"], df["split"] = None, "Unscored", "n/a"
    if segments is not None:
        df = df.merge(segments[["customer_id", "segment_name"]], on="customer_id", how="left")
    else:
        df["segment_name"] = "n/a"
    return df[DATA_COLS]


class Refs:
    """Absolute column ranges on the Data sheet, e.g. refs['churn'] -> Data!$N$2:$N$7044."""

    def __init__(self, cols: list[str], n: int):
        self.map = {c: f"Data!${get_column_letter(i + 1)}$2:${get_column_letter(i + 1)}${n + 1}"
                    for i, c in enumerate(cols)}

    def __getitem__(self, key):
        return self.map[key]


def _data_sheet(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Data")
    ws.append(DATA_COLS)
    for row in df.itertuples(index=False):
        ws.append(list(row))
    ref = f"A1:{get_column_letter(len(DATA_COLS))}{len(df) + 1}"
    tbl = Table(displayName="tblCustomers", ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tbl)
    ws.freeze_panes = "B2"
    formats = {"monthly_charges": MONEY2, "total_charges": MONEY2, "churn_probability": "0.000"}
    for i, col in enumerate(DATA_COLS, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(11, min(28, len(col) + 4))
        if col in formats:
            for cell in ws[letter][1:]:
                cell.number_format = formats[col]


def _assumptions_sheet(wb: Workbook, econ: dict, threshold: float) -> dict[str, str]:
    ws = wb.create_sheet("Assumptions")
    _title(ws, "Campaign assumptions", "Yellow cells are inputs - everything else recalculates from them")
    rows = [
        ("Offer cost per contacted customer ($)", econ["offer_cost"], MONEY2, "credit + agent handling time"),
        ("Months of revenue kept per saved churner", econ["months_of_value_saved"], "0", "avg. extra lifetime after accepting"),
        ("Offer acceptance rate among churners", econ["offer_success_rate"], PCT, "assumption - replace with the measured rate once a pilot has run"),
        ("Contact threshold (churn probability)", threshold, "0.00", "chosen on the validation set by the Python pipeline"),
    ]
    _header(ws, 4, ["Assumption", "Value", "Note"])
    names = {}
    for i, (label, val, fmt, note) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=label).font = _font()
        c = ws.cell(row=i, column=2, value=val)
        c.font = _font(color="0000FF")
        c.fill = INPUT_FILL
        c.number_format = fmt
        c.border = BOX
        ws.cell(row=i, column=3, value=note).font = _font(color="52514E", italic=True)
        names[label] = f"Assumptions!$B${i}"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 58
    return {
        "offer_cost": names[rows[0][0]],
        "months": names[rows[1][0]],
        "success": names[rows[2][0]],
        "threshold": names[rows[3][0]],
    }


def _breakdown(ws, top: int, title: str, field: str, levels: list[str], r: Refs) -> tuple[int, int]:
    ws.cell(row=top, column=1, value=title).font = _font(bold=True, size=12)
    _header(ws, top + 1, [field.replace("_", " ").title(), "Customers", "Churned", "Churn rate",
                          "MRR", "MRR lost", "Avg tenure (m)", "Share of base"])
    first = top + 2
    for i, lvl in enumerate(levels):
        row = first + i
        a = f"$A{row}"
        ws.cell(row=row, column=1, value=lvl)
        ws.cell(row=row, column=2, value=f"=COUNTIFS({r[field]},{a})")
        ws.cell(row=row, column=3, value=f"=COUNTIFS({r[field]},{a},{r['churn']},1)")
        ws.cell(row=row, column=4, value=f"=IFERROR(C{row}/B{row},0)")
        ws.cell(row=row, column=5, value=f"=SUMIFS({r['monthly_charges']},{r[field]},{a})")
        ws.cell(row=row, column=6, value=f"=SUMIFS({r['monthly_charges']},{r[field]},{a},{r['churn']},1)")
        ws.cell(row=row, column=7, value=f"=IFERROR(AVERAGEIFS({r['tenure']},{r[field]},{a}),0)")
        ws.cell(row=row, column=8, value=f"=IFERROR(B{row}/SUM($B${first}:$B${first + len(levels) - 1}),0)")
        for col, fmt in zip(range(1, 9), [None, INT, INT, PCT, MONEY, MONEY, "0.0", PCT]):
            c = ws.cell(row=row, column=col)
            c.font = _font()
            c.border = BOX
            if fmt:
                c.number_format = fmt
    last = first + len(levels) - 1
    ws.conditional_formatting.add(
        f"D{first}:D{last}",
        ColorScaleRule(start_type="min", start_color="FDF0EA", end_type="max", end_color="EB6834"),
    )
    return first, last


def _bar(ws, title: str, cats_ref, vals_ref, anchor: str, pct: bool = True) -> None:
    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = title
    ch.legend = None
    ch.add_data(vals_ref, titles_from_data=False)
    ch.set_categories(cats_ref)
    ch.series[0].graphicalProperties.solidFill = "EB6834"
    ch.series[0].graphicalProperties.line.noFill = True
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.numFmt = "0%" if pct else "#,##0"
    ch.y_axis.numFmt = "0%" if pct else "#,##0"
    ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.gapWidth = 60
    ch.height, ch.width = 6.5, 13
    ws.add_chart(ch, anchor)


def _summary_sheet(wb: Workbook, df: pd.DataFrame, r: Refs) -> None:
    ws = wb.create_sheet("Summary", 0)
    _title(ws, "Telco churn - executive summary", "IBM Telco Customer Churn sample, one snapshot of 7k customers")

    kpis = [
        ("Customers", f"=COUNTA({r['customer_id']})", INT),
        ("Churned", f"=SUM({r['churn']})", INT),
        ("Churn rate", f"=AVERAGE({r['churn']})", PCT),
        ("Monthly recurring revenue", f"=SUM({r['monthly_charges']})", MONEY),
        ("MRR lost to churn", f"=SUMIFS({r['monthly_charges']},{r['churn']},1)", MONEY),
        ("Share of MRR lost", "=IFERROR(B8/B7,0)", PCT),
        ("Avg tenure, churned (m)", f"=AVERAGEIFS({r['tenure']},{r['churn']},1)", "0.0"),
        ("Avg tenure, retained (m)", f"=AVERAGEIFS({r['tenure']},{r['churn']},0)", "0.0"),
    ]
    _header(ws, 3, ["KPI", "Value"])
    for i, (label, formula, fmt) in enumerate(kpis, start=4):
        ws.cell(row=i, column=1, value=label).font = _font(bold=True)
        c = ws.cell(row=i, column=2, value=formula)
        c.number_format, c.font, c.fill, c.border = fmt, _font(size=11, bold=True), KPI_FILL, BOX
        ws.cell(row=i, column=1).border = BOX

    top = 14
    sections = [
        ("By contract", "contract_type", ["Month-to-month", "One year", "Two year"]),
        ("By internet service", "internet_service", ["Fiber optic", "DSL", "No"]),
        ("By payment method", "payment_method", sorted(df["payment_method"].unique())),
        ("By tenure band", "tenure_band", TENURE_LABELS),
    ]
    for title, field, levels in sections:
        first, last = _breakdown(ws, top, title, field, levels, r)
        _bar(ws, f"Churn rate - {title.lower()[3:]}",
             Reference(ws, min_col=1, min_row=first, max_row=last),
             Reference(ws, min_col=4, min_row=first, max_row=last),
             f"J{top}")
        top = max(last + 3, top + 15)

    widths = [30, 13, 11, 11, 13, 13, 14, 13]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"


def _risk_sheet(wb: Workbook, r: Refs, a: dict) -> None:
    ws = wb.create_sheet("Risk & Campaign")
    _title(ws, "Churn risk and retention campaign", "Scores from the best model; change Assumptions to rerun the maths")

    _header(ws, 4, ["Risk tier", "Customers", "Avg predicted", "Actual churn rate",
                    "MRR in tier", "Expected MRR at risk"])
    for i, tier in enumerate(["Critical", "High", "Medium", "Low"], start=5):
        t = f"$A{i}"
        ws.cell(row=i, column=1, value=tier)
        ws.cell(row=i, column=2, value=f"=COUNTIFS({r['risk_tier']},{t})")
        ws.cell(row=i, column=3, value=f"=IFERROR(AVERAGEIFS({r['churn_probability']},{r['risk_tier']},{t}),0)")
        ws.cell(row=i, column=4, value=f"=IFERROR(COUNTIFS({r['risk_tier']},{t},{r['churn']},1)/B{i},0)")
        ws.cell(row=i, column=5, value=f"=SUMIFS({r['monthly_charges']},{r['risk_tier']},{t})")
        ws.cell(row=i, column=6, value=f"=SUMPRODUCT(({r['risk_tier']}={t})*{r['churn_probability']}*{r['monthly_charges']})")
        for col, fmt in zip(range(1, 7), [None, INT, PCT, PCT, MONEY, MONEY]):
            c = ws.cell(row=i, column=col)
            c.font, c.border = _font(), BOX
            if fmt:
                c.number_format = fmt
    ws.conditional_formatting.add("F5:F8", DataBarRule(start_type="min", end_type="max", color="EB6834"))

    ws["A11"] = "Campaign simulator (test split only, so results are out-of-sample)"
    ws["A11"].font = _font(bold=True, size=12)
    test = '"test"'
    thr = a["threshold"]
    tgt = f"{r['split']},{test},{r['churn_probability']},\">=\"&{thr}"
    lines = [
        ("Test customers", f"=COUNTIFS({r['split']},{test})", INT),
        ("Contacted (p >= threshold)", f"=COUNTIFS({tgt})", INT),
        ("  ...of whom actually churned", f"=COUNTIFS({tgt},{r['churn']},1)", INT),
        ("Churners in test split", f"=COUNTIFS({r['split']},{test},{r['churn']},1)", INT),
        ("Recall (churners reached)", "=IFERROR(B15/B16,0)", PCT),
        ("Precision (hit rate)", "=IFERROR(B15/B14,0)", PCT),
        ("Campaign cost", f"=B14*{a['offer_cost']}", MONEY),
        ("Revenue retained",
         f"={a['success']}*{a['months']}*SUMIFS({r['monthly_charges']},{tgt},{r['churn']},1)", MONEY),
        ("Net value", "=B20-B19", MONEY),
        ("ROI", "=IFERROR(B21/B19,0)", PCT),
    ]
    for i, (label, formula, fmt) in enumerate(lines, start=13):
        ws.cell(row=i, column=1, value=label).font = _font(bold=label in {"Net value", "ROI"})
        c = ws.cell(row=i, column=2, value=formula)
        c.number_format, c.border = fmt, BOX
        c.font = _font(bold=label in {"Net value", "ROI"})
        if label in {"Net value", "ROI"}:
            c.fill = KPI_FILL
    ws.column_dimensions["A"].width = 36
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18


def _segment_sheet(wb: Workbook, df: pd.DataFrame, r: Refs) -> None:
    ws = wb.create_sheet("Segments")
    _title(ws, "Customer segments (k-means)", "Clusters built without the churn label; churn only used to describe them")
    _header(ws, 4, ["Segment", "Customers", "Churn rate", "Avg monthly charge", "Avg tenure (m)", "Avg services", "MRR lost"])
    segs = sorted(df["segment_name"].dropna().unique())
    for i, seg in enumerate(segs, start=5):
        s = f"$A{i}"
        ws.cell(row=i, column=1, value=seg)
        ws.cell(row=i, column=2, value=f"=COUNTIFS({r['segment_name']},{s})")
        ws.cell(row=i, column=3, value=f"=IFERROR(COUNTIFS({r['segment_name']},{s},{r['churn']},1)/B{i},0)")
        ws.cell(row=i, column=4, value=f"=IFERROR(AVERAGEIFS({r['monthly_charges']},{r['segment_name']},{s}),0)")
        ws.cell(row=i, column=5, value=f"=IFERROR(AVERAGEIFS({r['tenure']},{r['segment_name']},{s}),0)")
        ws.cell(row=i, column=6, value=f"=IFERROR(AVERAGEIFS({r['n_services']},{r['segment_name']},{s}),0)")
        ws.cell(row=i, column=7, value=f"=SUMIFS({r['monthly_charges']},{r['segment_name']},{s},{r['churn']},1)")
        for col, fmt in zip(range(1, 8), [None, INT, PCT, MONEY2, "0.0", "0.0", MONEY]):
            c = ws.cell(row=i, column=col)
            c.font, c.border = _font(), BOX
            if fmt:
                c.number_format = fmt
    last = 4 + len(segs)
    if segs:
        ws.conditional_formatting.add(
            f"C5:C{last}", ColorScaleRule(start_type="min", start_color="FDF0EA", end_type="max", end_color="EB6834"))
        _bar(ws, "Churn rate by segment", Reference(ws, min_col=1, min_row=5, max_row=last),
             Reference(ws, min_col=3, min_row=5, max_row=last), f"A{last + 3}")
    ws.column_dimensions["A"].width = 38
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 16


def build(segments: pd.DataFrame | None = None, threshold: float = 0.5, out: Path | None = None) -> Path:
    cfg = load_config()
    out = out or ROOT / cfg["paths"]["excel"]
    out.parent.mkdir(parents=True, exist_ok=True)

    df = _load_frame(segments)
    r = Refs(DATA_COLS, len(df))

    wb = Workbook()
    wb.remove(wb.active)
    _data_sheet(wb, df)
    a = _assumptions_sheet(wb, cfg["economics"], threshold)
    _summary_sheet(wb, df, r)
    _risk_sheet(wb, r, a)
    _segment_sheet(wb, df, r)
    order = ["Summary", "Risk & Campaign", "Segments", "Assumptions", "Data"]
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(out)
    log.info("Excel report written to %s", out.relative_to(ROOT))
    return out
